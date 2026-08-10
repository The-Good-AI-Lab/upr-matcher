from __future__ import annotations

import re
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from fmsi_un_recommendations.utils import docx_tables_to_json, read_text_file

OLE_COMPOUND_DOCUMENT_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MATERIALS_ROOT = Path(
    "/home/arthur/Documents/gail/UPR Materials - TheGoodAILab-20260615T015654Z-3-001/UPR Materials - TheGoodAILab"
)
SOUTH_AFRICA_DOC = MATERIALS_ROOT / "Matrix of recommendations_SouthAfrica_2022.doc"
TANZANIA_DOCX = MATERIALS_ROOT / "Matrix of recommendations_Tanzania_2016.docx"
AUSTRALIA_DOCX = MATERIALS_ROOT / "Matrix of recommendations_Australia_2020.docx"
LEADING_TARGET_ID_RE = re.compile(r"(?m)^\s*(\d{2,3}\.\d{1,3})(?=\D)")


def _target_ids_in_text(text: str) -> set[str]:
    return set(LEADING_TARGET_ID_RE.findall(text))


def _target_ids_in_rows(rows: list[dict[str, str]]) -> set[str]:
    ids: set[str] = set()
    for row in rows:
        for value in row.values():
            match = LEADING_TARGET_ID_RE.search(value)
            if match:
                ids.add(match.group(1))
                break
    return ids


def _row_id_coverage(path: Path) -> float:
    text_ids = _target_ids_in_text(read_text_file(path))
    row_ids = _target_ids_in_rows(docx_tables_to_json(path))
    if not text_ids:
        return 0.0
    return len(text_ids & row_ids) / len(text_ids)


class WordDocumentParsingTests(unittest.TestCase):
    def test_doc_suffix_with_docx_package_can_be_read(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "matrix.doc"
            document = Document()
            document.add_paragraph("Synthetic UPR matrix")
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Recommendation and recommending State"
            table.rows[0].cells[1].text = "Position"
            table.rows[1].cells[0].text = "120.1 Protect children's rights (State);"
            table.rows[1].cells[1].text = "Supported"
            document.save(path)

            text = read_text_file(path)
            rows = docx_tables_to_json(path)

        self.assertIn("120.1 Protect children's rights", text)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Position"], "Supported")

    def test_theme_first_table_without_header_preserves_recommendation(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "theme-first.docx"
            document = Document()
            table = document.add_table(rows=2, cols=4)
            for cell in table.rows[0].cells:
                cell.text = "Theme: A47 Good governance"
            table.rows[1].cells[0].text = "134.128 Finalize the anti-corruption action plan (Morocco);"
            table.rows[1].cells[1].text = "Morocco"
            table.rows[1].cells[2].text = "Supported"
            table.rows[1].cells[3].text = "UNDP"
            document.save(path)

            rows = docx_tables_to_json(path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Theme"], "A47 Good governance")
        self.assertIn("134.128", rows[0]["Recommendation"])

    def test_merged_header_table_preserves_first_duplicate_column(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "merged-header.docx"
            document = Document()
            table = document.add_table(rows=2, cols=3)
            table.rows[0].cells[0].merge(table.rows[0].cells[1])
            table.rows[0].cells[0].text = "Recommendation"
            table.rows[0].cells[2].text = "Position"
            table.rows[1].cells[0].text = "146.1 Ratify the treaty (State);"
            table.rows[1].cells[2].text = "Supported"
            document.save(path)

            rows = docx_tables_to_json(path)

        self.assertEqual(len(rows), 1)
        self.assertIn("146.1", rows[0]["Recommendation"])
        self.assertEqual(rows[0]["Position"], "Supported")

    def test_right_or_area_metadata_row_is_not_a_recommendation(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "right-or-area.docx"
            document = Document()
            table = document.add_table(rows=3, cols=3)
            table.rows[0].cells[0].text = "Recommendation"
            table.rows[0].cells[1].text = "Position"
            table.rows[0].cells[2].text = "Themes"
            for cell in table.rows[1].cells:
                cell.text = "Right or area: 2.1 Acceptance of international norms"
            table.rows[2].cells[0].text = "130.1 Ratify the OP-CAT (Palestine);"
            table.rows[2].cells[1].text = "Supported"
            table.rows[2].cells[2].text = "2.1 Acceptance of international norms"
            document.save(path)

            rows = docx_tables_to_json(path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Right or area"], "2.1 Acceptance of international norms")
        self.assertIn("130.1", rows[0]["Recommendation"])

    def test_themes_header_is_not_treated_as_metadata(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "themes-header.docx"
            document = Document()
            table = document.add_table(rows=2, cols=3)
            table.rows[0].cells[0].text = "Themes"
            table.rows[0].cells[1].text = "Recommendation"
            table.rows[0].cells[2].text = "Position"
            table.rows[1].cells[0].text = "A12 Freedom of expression"
            table.rows[1].cells[1].text = "140.1 Protect journalists (State);"
            table.rows[1].cells[2].text = "Supported"
            document.save(path)

            rows = docx_tables_to_json(path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Themes"], "A12 Freedom of expression")
        self.assertIn("140.1", rows[0]["Recommendation"])
        self.assertEqual(rows[0]["Position"], "Supported")

    def test_legacy_doc_has_actionable_error(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "legacy.doc"
            path.write_bytes(OLE_COMPOUND_DOCUMENT_MAGIC + b"\x00" * 128)

            with self.assertRaisesRegex(ValueError, "Legacy \\.doc files are not supported.*Convert"):
                read_text_file(path)
            with self.assertRaisesRegex(ValueError, "Legacy \\.doc files are not supported.*Convert"):
                docx_tables_to_json(path)

    @unittest.skipUnless(SOUTH_AFRICA_DOC.exists(), "External UPR materials fixture not present")
    def test_real_doc_extension_docx_package_is_supported(self) -> None:
        text = read_text_file(SOUTH_AFRICA_DOC)
        rows = docx_tables_to_json(SOUTH_AFRICA_DOC)

        self.assertGreater(len(text), 100_000)
        self.assertGreater(len(rows), 200)

    @unittest.skipUnless(TANZANIA_DOCX.exists(), "External UPR materials fixture not present")
    def test_tanzania_fixture_preserves_target_ids(self) -> None:
        self.assertGreaterEqual(_row_id_coverage(TANZANIA_DOCX), 0.98)

    @unittest.skipUnless(AUSTRALIA_DOCX.exists(), "External UPR materials fixture not present")
    def test_australia_fixture_preserves_target_ids(self) -> None:
        self.assertGreaterEqual(_row_id_coverage(AUSTRALIA_DOCX), 0.98)
